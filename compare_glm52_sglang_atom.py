#!/usr/bin/env python3
"""
compare_glm52_sglang_atom.py

GLM-5.2-specific, apples-to-apples kernel comparison between an SGLang trace and
an ATOM (rocm/atom-dev) trace for ONE forward pass (prefill chunk or decode step).

Why a bespoke tool: analyze_trace.py (SGLang) and analyze_atom_trace.py (ATOM)
emit different Section/Kernel taxonomies, so compare_breakdown.py's row alignment
falls apart. This script instead maps BOTH stacks' kernels into the SAME GLM-5.2
functional buckets (sparse-MLA attn, DSA indexer+topk, MLA/dense GEMM, MoE
up/gate GEMM, MoE down GEMM, all-reduce, rmsnorm/quant, rope/kv-cache, other),
so the two columns are directly comparable.

It is intentionally GLM-5.2-only: bucket regexes and the 16384-token prefill /
bs=64 decode heuristics assume the GLM-5.2 DSA+MLA+MoE architecture.

Segmentation is stream-aware (kernels + wrapper annotations restricted to the
dominant GPU stream), isolating ONE forward between consecutive wrapper markers:
  - SGLang:  gpu_user_annotation  "step[EXTEND ...]" / "step[DECODE ...]"
  - ATOM:    gpu_user_annotation  "prefill[...]"     / "decode[...]"

For SGLang prefill, triton kernels can be JIT/autotune-inflated on their first
occurrence; when several qualifying forwards exist we take the FASTEST (min
Sigma) as the steady-state estimate and warn if only one is available.

Usage:
  python compare_glm52_sglang_atom.py --phase prefill \
      --sglang TRACE_SGLANG_EXTEND.json.gz --atom TRACE_ATOM.json.gz \
      --out cmp_prefill.csv
  python compare_glm52_sglang_atom.py --phase decode \
      --sglang TRACE_SGLANG_DECODE.json.gz --atom TRACE_ATOM.json.gz \
      --out cmp_decode.csv
"""
from __future__ import annotations

import argparse
import bisect
import gzip
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path


# --------------------------------------------------------------------------- #
# GLM-5.2 functional buckets. Order = display order. First matching regex wins.
# Patterns are lowercased substring/regex tests applied to the kernel name.
# --------------------------------------------------------------------------- #
import re

BUCKET_RULES = [
    # all-reduce / TP communication (quickreduce INT4, rccl, aiter allgather, etc.)
    ("all-reduce/comm", [
        r"reduce_scatter", r"all_?reduce", r"allgather", r"all_gather",
        r"nccl", r"rccl", r"cross_device", r"quick.*reduce", r"custom_all",
    ]),
    # DSA indexer + top-k selection (paged MQA logits, hadamard, fp8 index quant).
    # NOTE: use SPECIFIC topk names (topk_transform / radix_topk) so we don't
    # swallow MoE's expert top-k kernels (grouped_topk, moe_reduction_..._topk9).
    ("DSA indexer+topk", [
        r"paged_mqa_logits", r"mqa_logits", r"deepgemm_fp8_paged", r"hadamard",
        r"topk_transform", r"radix_topk", r"indexer", r"transform_index",
        r"convert_req_index", r"fused_qk_rmsnorm_group_quant",
    ]),
    # sparse MLA attention core (triton sparse-mla, aiter mla, tilelang mla)
    ("sparse-MLA attn", [
        r"sparse_mla", r"_mla_", r"\bmla\b", r"mla_decode", r"mla_fwd",
        r"flash.*mla", r"aiter\d*::mla", r"mla_a8w8",
    ]),
    # MoE experts up/gate projection GEMM (moe1 / silu-mul fused)
    ("MoE up/gate GEMM (moe1)", [
        r"moe1", r"moe_?stage1", r"gate_?up", r"silu_mul", r"silu_and_mul",
    ]),
    # MoE experts down projection GEMM (moe2)
    ("MoE down GEMM (moe2)", [
        r"moe2", r"moe_?stage2", r"down_proj_moe",
    ]),
    # other MoE plumbing: routing/sorting, expert top-k, output reduction, gather
    ("MoE routing/other", [
        r"moe_sort", r"moe_align", r"moe_sorting", r"fused_moe", r"fused_mx_quant_moe",
        r"grouped_?gemm", r"group_?gemm", r"grouped_topk", r"moe_reduction",
        r"append_shared_expert", r"\bmoe\b", r"expert",
    ]),
    # dense / linear GEMM: attention q/kv/o projections + dense MLP (non-MoE)
    ("dense/linear GEMM", [
        r"cijk", r"hgemm", r"\bgemm\b", r"cshuffle", r"gemm_xdl", r"tensile",
        r"matmul", r"wgrad", r"a8w8", r"f8_.*gemm", r"gemm_a8", r"gemm_afp4",
    ]),
    # rmsnorm / quant / activation (non-indexer)
    ("rmsnorm/quant/act", [
        r"rmsnorm", r"\bnorm\b", r"layernorm", r"quant", r"dequant", r"silu",
        r"gelu", r"activation", r"scaled_",
    ]),
    # rope + kv-cache writes
    ("rope/kv-cache", [
        r"rope", r"kv_?cache", r"cache_flat", r"set_kv", r"store_kv",
        r"reshape_and_cache", r"append_kv",
    ]),
    # embedding
    ("embedding", [r"embed"]),
    # elementwise plumbing
    ("elementwise/other", [
        r"elementwise", r"\bcopy\b", r"copybuffer", r"\bcast\b", r"\bfill\b",
        r"memset", r"\badd\b", r"vectorized_elementwise", r"index_", r"permute",
        r"transpose", r"\bcat\b", r"concat", r"masked_embedding",
    ]),
]
COMPILED = [(name, [re.compile(p) for p in pats]) for name, pats in BUCKET_RULES]


def classify(kernel_name: str) -> str:
    n = kernel_name.lower()
    for name, pats in COMPILED:
        for p in pats:
            if p.search(n):
                return name
    return "other"


# --------------------------------------------------------------------------- #
def load_trace(path: str):
    p = Path(path)
    if not p.exists():
        sys.exit(f"[ERROR] file not found: {path}")
    opener = gzip.open if path.endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8") as f:
        t = json.load(f)
    return t if isinstance(t, list) else t.get("traceEvents", [])


def _x(e):
    return isinstance(e, dict) and e.get("ph") == "X"


def dominant_stream(kernels):
    return Counter((k["pid"], k["tid"]) for k in kernels).most_common(1)[0][0]


def collect(evs):
    kernels = [e for e in evs if _x(e) and str(e.get("cat", "")).lower() == "kernel"]
    if not kernels:
        sys.exit("[ERROR] no kernel events in trace")
    dom = dominant_stream(kernels)
    kernels = [{"name": k.get("name", "?"), "ts": float(k["ts"]),
                "dur": float(k.get("dur", 0.0))}
               for k in kernels if (k["pid"], k["tid"]) == dom]
    kernels.sort(key=lambda k: k["ts"])
    ann = [{"name": a.get("name", "?"), "ts": float(a["ts"]), "dur": float(a.get("dur", 0.0))}
           for a in evs if _x(a) and a.get("cat") == "gpu_user_annotation"
           and (a["pid"], a["tid"]) == dom]
    ann.sort(key=lambda a: a["ts"])
    return kernels, ann, dom


def wrapper_markers(ann, phase, stack):
    """Return the forward-boundary wrapper annotations (all forwards), sorted."""
    if stack == "sglang":
        prefixes = ("step[",)
    else:  # atom
        prefixes = ("prefill[", "decode[")
    return [a for a in ann if any(a["name"].startswith(p) for p in prefixes)]


def _is_target(name, phase, stack):
    n = name
    if phase == "prefill":
        if stack == "sglang":
            return n.startswith("step[EXTEND") and "toks=16" in n
        return n.startswith("prefill[") and "tok=16384" in n
    else:  # decode
        if stack == "sglang":
            return n.startswith("step[DECODE") or (n.startswith("step[") and "DECODE" in n)
        return n.startswith("decode[") and "bs=64" in n


def segment_and_sum(kernels, ann, phase, stack):
    """Isolate one representative forward and return (label, span_us, {bucket: us},
    total_us, nkernels, per_kernel[(bucket,name,us,count)])."""
    marks = wrapper_markers(ann, phase, stack)
    if not marks:
        sys.exit(f"[ERROR] no {stack} {phase} wrapper markers found on dominant stream")
    mts = [m["ts"] for m in marks]
    kts = [k["ts"] for k in kernels]

    def window_sum(s0, s1):
        lo = bisect.bisect_left(kts, s0)
        hi = bisect.bisect_left(kts, s1)
        return kernels[lo:hi]

    # candidate target forwards
    cands = []
    for m in marks:
        if not _is_target(m["name"], phase, stack):
            continue
        s0 = m["ts"]
        j = bisect.bisect_right(mts, s0)
        s1 = mts[j] if j < len(mts) else (kts[-1] + kernels[-1]["dur"] + 1)
        win = window_sum(s0, s1)
        tot = sum(k["dur"] for k in win)
        cands.append((tot, m["name"], s1 - s0, win))
    if not cands:
        sys.exit(f"[ERROR] no target {phase} forward (full-chunk/bs=64) found for {stack}")

    # SGLang triton prefill may be JIT-inflated on first occurrence -> take the
    # fastest (min Sigma) as steady-state; decode/ATOM -> median is fine.
    if stack == "sglang" and phase == "prefill" and len(cands) > 1:
        cands.sort(key=lambda c: c[0])
        chosen = cands[0]
        note = f"min-of-{len(cands)} (JIT-guard)"
    else:
        cands.sort(key=lambda c: c[0])
        chosen = cands[len(cands) // 2]
        note = f"median-of-{len(cands)}" if len(cands) > 1 else "only-1"

    tot, label, span, win = chosen
    buckets = defaultdict(float)
    perk = defaultdict(lambda: [0.0, 0])
    for k in win:
        b = classify(k["name"])
        buckets[b] += k["dur"]
        perk[(b, k["name"])][0] += k["dur"]
        perk[(b, k["name"])][1] += 1
    perk_rows = [(b, nm, us, c) for (b, nm), (us, c) in perk.items()]
    return label, span, dict(buckets), tot, len(win), perk_rows, note


BUCKET_ORDER = [name for name, _ in BUCKET_RULES] + ["other"]


def write_grouped_xlsx(path, phase, labels, meta, cats, sb, ab, sperk, aperk):
    """Pretty module-grouped side-by-side workbook (compare_breakdown.py style).

    One block per functional module (bucket): a filled header row with each
    side's Σ / % / count, then the kernels of both sides listed side by side
    (rank-aligned by descending duration, shorter side padded)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    LA, LB = labels
    stot = sum(sb.values()) or 1.0
    atot = sum(ab.values()) or 1.0

    # per-bucket kernel lists: [(sum_ms, avg_us, count, name), ...] desc
    def by_bucket(perk):
        d = {}
        for b, nm, us, c in perk:
            d.setdefault(b, []).append((us / 1000.0, us / c if c else 0.0, c, nm))
        for b in d:
            d[b].sort(reverse=True)
        return d
    sK, aK = by_bucket(sperk), by_bucket(aperk)

    wb = Workbook(); ws = wb.active; ws.title = f"{phase}_by_module"
    bold = Font(name="Arial", size=10, bold=True)
    reg = Font(name="Arial", size=9)
    mono = Font(name="Consolas", size=9)
    hdr_fill = PatternFill("solid", fgColor="305496")
    mod_fill = PatternFill("solid", fgColor="D9E1F2")
    tot_fill = PatternFill("solid", fgColor="FCE4D6")
    white = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center"); left = Alignment(horizontal="left")

    headers = ["Module", f"{LA} Kernel", "Avg_us", "Σ_ms", "Cnt", "",
               f"{LB} Kernel", "Avg_us", "Σ_ms", "Cnt"]
    # title + meta
    ws.cell(1, 1, f"GLM-5.2 {phase} — {LA} vs {LB} (per forward, grouped by module)").font = bold
    ws.cell(2, 1, meta[0]).font = reg
    ws.cell(3, 1, meta[1]).font = reg
    r = 5
    for c, h in enumerate(headers, 1):
        cell = ws.cell(r, c, h); cell.font = white if h else reg
        cell.fill = hdr_fill if h else PatternFill(); cell.alignment = center
    ws.freeze_panes = f"A{r+1}"
    r += 1

    for b in cats:
        s_ms = sb.get(b, 0.0) / 1000.0
        a_ms = ab.get(b, 0.0) / 1000.0
        ratio = (s_ms / a_ms) if a_ms > 1e-9 else float("inf")
        # module header row
        ws.cell(r, 1, b).font = bold
        ws.cell(r, 4, round(s_ms, 3)).font = bold
        ws.cell(r, 3, f"{s_ms/ (stot/1000):.0%}").font = reg
        ws.cell(r, 9, round(a_ms, 3)).font = bold
        ws.cell(r, 8, f"{a_ms/ (atot/1000):.0%}").font = reg
        ws.cell(r, 5, f"A/B={ratio:.2f}").font = bold
        for c in range(1, 11):
            ws.cell(r, c).fill = mod_fill; ws.cell(r, c).border = border
        r += 1
        # kernel rows
        sk = sK.get(b, []); ak = aK.get(b, [])
        for i in range(max(len(sk), len(ak))):
            if i < len(sk):
                sm, sa, sc, snm = sk[i]
                ws.cell(r, 2, snm[:70]).font = mono
                ws.cell(r, 3, round(sa, 2)).font = reg
                ws.cell(r, 4, round(sm, 3)).font = reg
                ws.cell(r, 5, sc).font = reg
            if i < len(ak):
                am, aa, ac, anm = ak[i]
                ws.cell(r, 7, anm[:70]).font = mono
                ws.cell(r, 8, round(aa, 2)).font = reg
                ws.cell(r, 9, round(am, 3)).font = reg
                ws.cell(r, 10, ac).font = reg
            r += 1

    # total row
    ws.cell(r, 1, "TOTAL").font = bold
    ws.cell(r, 4, round(stot / 1000, 3)).font = bold
    ws.cell(r, 9, round(atot / 1000, 3)).font = bold
    ws.cell(r, 5, f"A/B={ (stot/atot) if atot else 0:.2f}").font = bold
    for c in range(1, 11):
        ws.cell(r, c).fill = tot_fill; ws.cell(r, c).border = border

    widths = [26, 60, 9, 9, 6, 3, 60, 9, 9, 6]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(path)
    print(f"[INFO] grouped xlsx written: {path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--phase", choices=["prefill", "decode"], required=True)
    ap.add_argument("--sglang", required=True, metavar="TRACE",
                    help="SGLang graph-ON trace: EXTEND for prefill, DECODE for decode")
    ap.add_argument("--atom", required=True, metavar="TRACE",
                    help="ATOM graph-ON trace (combined)")
    ap.add_argument("--out", required=True, metavar="CSV")
    ap.add_argument("--xlsx", metavar="XLSX", default=None,
                    help="also write a pretty module-grouped side-by-side workbook")
    ap.add_argument("--labels", nargs=2, default=["SGLANG", "ATOM"])
    ap.add_argument("--top", type=int, default=6,
                    help="top-N kernels per bucket to include in the detail CSV")
    args = ap.parse_args()

    sk, sa, sd = collect(load_trace(args.sglang))
    ak, aa, ad = collect(load_trace(args.atom))
    slabel, sspan, sb, stot, snk, sperk, snote = segment_and_sum(sk, sa, args.phase, "sglang")
    alabel, aspan, ab, atot, ank, aperk, anote = segment_and_sum(ak, aa, args.phase, "atom")

    LA, LB = args.labels
    print(f"\n=== GLM-5.2 {args.phase} comparison ({LA} vs {LB}) ===")
    print(f"  {LA}: {slabel}")
    print(f"        forward={snote}  span={sspan/1000:.2f}ms  Sigma_kernel={stot/1000:.2f}ms  nkernels={snk}")
    print(f"  {LB}: {alabel}")
    print(f"        forward={anote}  span={aspan/1000:.2f}ms  Sigma_kernel={atot/1000:.2f}ms  nkernels={ank}")
    cats = sorted(set(sb) | set(ab), key=lambda c: BUCKET_ORDER.index(c) if c in BUCKET_ORDER else 99)
    print(f"\n  {'Bucket':<24}{LA+'_ms':>12}{LB+'_ms':>12}{'Δ(A-B)ms':>12}{'A/B':>8}")
    print(f"  {'-'*24}{'-'*12}{'-'*12}{'-'*12}{'-'*8}")
    for c in cats:
        a = sb.get(c, 0.0) / 1000
        b = ab.get(c, 0.0) / 1000
        ratio = (a / b) if b > 1e-9 else float("inf")
        print(f"  {c:<24}{a:>12.2f}{b:>12.2f}{a-b:>12.2f}{ratio:>8.2f}")
    print(f"  {'-'*24}{'-'*12}{'-'*12}{'-'*12}{'-'*8}")
    print(f"  {'TOTAL Sigma_kernel':<24}{stot/1000:>12.2f}{atot/1000:>12.2f}{(stot-atot)/1000:>12.2f}"
          f"{(stot/atot if atot else float('inf')):>8.2f}")

    # --- write CSV -------------------------------------------------------- #
    import csv
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([f"# GLM-5.2 {args.phase} comparison", LA, "vs", LB])
        w.writerow([f"# {LA} forward", slabel, snote, f"span_ms={sspan/1000:.2f}",
                    f"sigma_ms={stot/1000:.2f}", f"nkernels={snk}"])
        w.writerow([f"# {LB} forward", alabel, anote, f"span_ms={aspan/1000:.2f}",
                    f"sigma_ms={atot/1000:.2f}", f"nkernels={ank}"])
        w.writerow([])
        w.writerow(["Bucket", f"{LA}_ms", f"{LB}_ms", "Delta_A_minus_B_ms", "A_over_B"])
        for c in cats:
            a = sb.get(c, 0.0) / 1000
            b = ab.get(c, 0.0) / 1000
            ratio = round(a / b, 3) if b > 1e-9 else ""
            w.writerow([c, round(a, 3), round(b, 3), round(a - b, 3), ratio])
        w.writerow(["TOTAL", round(stot/1000, 3), round(atot/1000, 3),
                    round((stot-atot)/1000, 3),
                    round(stot/atot, 3) if atot else ""])
        # per-bucket top kernels for both sides
        w.writerow([])
        w.writerow([f"# top-{args.top} kernels per bucket"])
        w.writerow(["Side", "Bucket", "Kernel", "Sum_ms", "Count"])
        for side, perk in ((LA, sperk), (LB, aperk)):
            by_bucket = defaultdict(list)
            for b, nm, us, c in perk:
                by_bucket[b].append((us, nm, c))
            for c in cats:
                for us, nm, cnt in sorted(by_bucket.get(c, []), reverse=True)[:args.top]:
                    w.writerow([side, c, nm[:80], round(us/1000, 3), cnt])
    print(f"\n[INFO] CSV written: {out}")

    if args.xlsx:
        meta = (f"{LA}: {slabel}  [{snote}, span={sspan/1000:.1f}ms, Σ={stot/1000:.1f}ms, {snk} kernels]",
                f"{LB}: {alabel}  [{anote}, span={aspan/1000:.1f}ms, Σ={atot/1000:.1f}ms, {ank} kernels]")
        write_grouped_xlsx(args.xlsx, args.phase, (LA, LB), meta, cats, sb, ab, sperk, aperk)


if __name__ == "__main__":
    main()
