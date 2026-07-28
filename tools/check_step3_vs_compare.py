#!/usr/bin/env python3
"""Sanity-check that a step3_layer_breakdown workbook is on the same per-forward
scale as a compare_glm52 bucket CSV: totals and the all-reduce row should agree
to within a few percent (step3 drops the handful of kernels that run outside the
decoder layers, so exact equality is not expected).

Usage: python check_step3_vs_compare.py --step3 A.xlsx [--step3 B.xlsx] --csv cmp.csv
"""
from __future__ import annotations

import argparse
import csv
import re

from openpyxl import load_workbook

AR = re.compile(r"allreduce_prototype|quickreduce|reduce_scatter", re.I)


def read(path):
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    h = {n: i for i, n in enumerate(rows[0]) if n}
    tot = ar = 0.0
    ar_calls = 0
    for r in rows[1:]:
        kn = r[h["KernelName"]]
        if kn is None:
            continue
        s = r[h["SumDuration_us"]] or 0
        if not isinstance(s, (int, float)):
            continue
        tot += s
        if AR.search(str(kn)):
            ar += s
            c = r[h["Count"]]
            ar_calls += c if isinstance(c, int) else 0
    return tot / 1000.0, ar / 1000.0, ar_calls


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--step3", action="append", required=True)
    ap.add_argument("--csv", required=True)
    args = ap.parse_args()

    want = {}
    with open(args.csv) as f:
        for r in csv.reader(f):
            if len(r) >= 3 and r[0] in ("all-reduce/comm", "TOTAL"):
                want[r[0]] = (r[1], r[2])
    print(f"compare_glm52 ({args.csv}):")
    for k, v in want.items():
        print(f"  {k:<18} colA={v[0]:>8}  colB={v[1]:>8}")
    print("step3 workbooks:")
    for p in args.step3:
        tot, ar, n = read(p)
        per = (ar * 1000 / n) if n else 0
        print(f"  {p.split('/')[-1]:<44} total={tot:8.2f} ms  "
              f"all-reduce={ar:7.2f} ms over {n} calls ({per:.1f} us/call)")


if __name__ == "__main__":
    main()
