#!/usr/bin/env python3
"""
glm52_buckets.py

GLM-5.2 functional-bucket comparison of two (or more) already-analyzed runs, and
the single source of truth for the bucket rules themselves (imported by
side_by_side.py).

It consumes the workbooks written by analyze/sglang_trace.py /
analyze/atom_trace.py instead of re-parsing the raw traces, so both sides are
compared over the SAME one forward those tools already isolated
(--forward-match / --forward-pick), with no second segmentation policy here:

  step3_layer_breakdown_*.xlsx -> decoder layers only  (default scope; ~99.8%
                                  of prefill, ~98.5% of decode kernel time)
  step1_kernel_stats_*.xlsx    -> the whole forward, lm_head / sampling / logits
                                  all-gather tail included

The scope is taken from whichever workbook you pass and is printed per side.
Mixing scopes across sides is refused.

Usage:
  python glm52_buckets.py --phase prefill --out cmp.csv \
      --src SGLANG step3_layer_breakdown_prefill_SGLANG.xlsx \
      --src ATOM   step3_layer_breakdown_prefill_ATOM.xlsx
"""
import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


# --------------------------------------------------------------------------- #
# GLM-5.2 functional buckets. Order = display order. First matching regex wins.
# Patterns are lowercased substring/regex tests applied to the kernel name.
# --------------------------------------------------------------------------- #
import re

BUCKET_RULES = [
    # all-reduce / TP communication (quickreduce INT4, rccl, aiter allgather, etc.)
    ("all-reduce/comm", [
        r"reduce_scatter", r"all_?reduce", r"allgather", r"all_gather",
        r"nccl", r"rccl", r"cross_device", r"quick.*reduce", r"custom_all",
    ]),
    # DSA indexer + top-k selection (paged MQA logits, hadamard, fp8 index quant).
    # NOTE: use SPECIFIC topk names (topk_transform / radix_topk) so we don't
    # swallow MoE's expert top-k kernels (grouped_topk, moe_reduction_..._topk9).
    ("DSA indexer+topk", [
        r"paged_mqa_logits", r"mqa_logits", r"deepgemm_fp8_paged", r"hadamard",
        r"topk_transform", r"radix_topk", r"indexer", r"transform_index",
        r"convert_req_index", r"fused_qk_rmsnorm_group_quant",
        r"dsa_decode_metadata",
        # sgl-kernel's topk family, which B200 uses for the indexer: the module
        # tree puts every one of them under DeepseekV2AttentionMLA > Indexer,
        # opposite MI355X's topk_transform_decode_kernel. Named specifically so
        # they still do not swallow MoE's grouped_topk.
        r"topk_main_kernel", r"topk_persistent_cluster", r"topk_small_batch",
        r"topk_plan",
    ]),
    # sparse MLA attention core (triton sparse-mla, aiter mla, tilelang mla).
    # B200/trtllm names it after the head dims instead: HQk576 (512 nope + 64
    # rope) and HV512 are MLA's, and TokenSparse is the DSA path. Confirmed by
    # the module tree -- it is the only kernel DeepseekV2AttentionMLA (self)
    # launches at that size.
    ("sparse-MLA attn", [
        r"sparse_mla", r"_mla_", r"\bmla\b", r"mla_decode", r"mla_fwd",
        r"flash.*mla", r"aiter\d*::mla", r"mla_a8w8",
        r"fmhasm100", r"hqk576",
    ]),
    # MoE experts up/gate projection GEMM (moe1 / silu-mul fused).
    # gemm1_a4w4_port_* / mxmoe_g1 are the FlyDSL MoE stage-1 kernels the aiter in
    # rocm/sgl-dev:v0.5.16-...-20260729 picks up from glm5_fp4_tuned_fmoe.csv; the
    # older mfma_moe1_* path is what images without that tuned CSV run.
    # bmm_E2m1_* is B200's: the two trtllm NVFP4 grouped GEMMs both sit under
    # DeepseekV2MoE (self) so the module tree cannot split them, but execution
    # order (Index 27 then 28) and output dtype do -- stage 1 emits E2m1 to feed
    # stage 2, stage 2 emits Bfloat16 into the residual.
    ("MoE up/gate GEMM (moe1)", [
        r"moe1", r"moe_?stage1", r"gate_?up", r"silu_mul", r"silu_and_mul",
        r"gemm1_a4w4", r"mxmoe_g1", r"bmm_e2m1_",
    ]),
    # MoE experts down projection GEMM (moe2)
    ("MoE down GEMM (moe2)", [
        r"moe2", r"moe_?stage2", r"down_proj_moe",
        r"gemm2_a4w4", r"mxmoe_g2", r"bmm_bfloat16_e2m1",
    ]),
    # other MoE plumbing: routing/sorting, expert top-k, output reduction, gather.
    # mxfp4_moe:: covers the FlyDSL path's separate scatter-reduce and 3-stage sort,
    # whose work the mfma moe2 kernel folds into its own epilogue.
    ("MoE routing/other", [
        r"moe_sort", r"moe_align", r"moe_sorting", r"fused_moe", r"fused_mx_quant_moe",
        r"grouped_?gemm", r"group_?gemm", r"grouped_topk", r"moe_reduction",
        r"append_shared_expert", r"mxfp4_moe", r"\bmoe\b", r"expert",
        # B200's router projection, under DeepseekV2MoE > MoEGate -- the same
        # place MI355X's grouped_topk_kernel sits, so routing on both sides.
        r"router_gemm",
    ]),
    # dense / linear GEMM: attention q/kv/o projections + dense MLP (non-MoE).
    # nvjet_sm100_* is cuBLAS on Blackwell; the module tree puts every one of
    # them under a Linear (attention projections, prepare_mlp, the MoE shared
    # expert), same call sites as MI355X's Cijk/bf16gemm. bf16gemm has no word
    # boundary before "gemm", so \bgemm\b missed it and 83 ms of MI355X
    # ColumnParallelLinear was landing in "other".
    ("dense/linear GEMM", [
        r"cijk", r"hgemm", r"\bgemm\b", r"cshuffle", r"gemm_xdl", r"tensile",
        r"matmul", r"wgrad", r"a8w8", r"f8_.*gemm", r"gemm_a8", r"gemm_afp4",
        r"bf16gemm", r"nvjet", r"cublas", r"splitkreduce", r"fused_a_gemm",
    ]),
    # rope + kv-cache writes. Ahead of rmsnorm/quant because B200 fuses the two:
    # flashinfer's RopeQuantizeKernel would otherwise be caught by "quant" and
    # counted against MI355X, whose rope (aiter kn_entry_2c_sbhd_cached, under
    # Indexer > RotaryEmbedding) is a separate kernel. No ROCm kernel matches
    # both sets, so the reorder leaves that side's classification unchanged.
    ("rope/kv-cache", [
        r"rope", r"kv_?cache", r"cache_flat", r"set_kv", r"store_kv",
        r"reshape_and_cache", r"append_kv", r"sbhd_cached",
    ]),
    # rmsnorm / quant / activation (non-indexer)
    ("rmsnorm/quant/act", [
        r"rmsnorm", r"\bnorm\b", r"layernorm", r"quant", r"dequant", r"silu",
        r"gelu", r"activation", r"scaled_",
    ]),
    # embedding
    ("embedding", [r"embed"]),
    # elementwise plumbing
    ("elementwise/other", [
        r"elementwise", r"\bcopy\b", r"copybuffer", r"\bcast\b", r"\bfill\b",
        r"memset", r"\badd\b", r"vectorized_elementwise", r"index_", r"permute",
        r"transpose", r"\bcat\b", r"concat", r"masked_embedding",
    ]),
]
COMPILED = [(name, [re.compile(p) for p in pats]) for name, pats in BUCKET_RULES]

def classify(kernel_name: str) -> str:
    n = kernel_name.lower()
    for name, pats in COMPILED:
        for p in pats:
            if p.search(n):
                return name
    return "other"


BUCKET_ORDER = [name for name, _ in BUCKET_RULES] + ["other"]


def read_workbook(path):
    """-> (scope, [(kernel_name, sum_us, launch_count)]).

    Accepts a step3 breakdown (one row per call site) or a step1 kernel-stats
    sheet (one row per kernel name); scope is detected from the header."""
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit(f"[ERROR] {path}: empty sheet")
    h = {n: i for i, n in enumerate(rows[0]) if n}

    def col(*names):
        for n in names:
            if n in h:
                return h[n]
        return None

    kn = col("KernelName", "Name")
    us = col("SumDuration_us")
    cnt = col("LaunchCount", "Count")
    if kn is None or us is None:
        raise SystemExit(f"[ERROR] {path}: not a step1/step3 workbook "
                         f"(header={list(h)[:8]})")
    scope = "layers" if "CallSite" in h or "LayerType" in h else "forward"
    out = []
    for r in rows[1:]:
        name = r[kn]
        if not name or str(name).startswith(("TOTAL", "Subtotal", "#")):
            continue
        out.append((str(name), float(r[us] or 0.0),
                    int(r[cnt] or 0) if cnt is not None else 0))
    return scope, out


def bucketize(rows):
    """-> ({bucket: sum_us}, [(bucket, kernel, sum_us, count)], total_us)."""
    buckets = defaultdict(float)
    perk = defaultdict(lambda: [0.0, 0])
    for name, sum_us, cnt in rows:
        b = classify(name)
        buckets[b] += sum_us
        e = perk[(b, name)]
        e[0] += sum_us
        e[1] += cnt
    perk_rows = [(b, nm, v[0], v[1]) for (b, nm), v in perk.items()]
    return dict(buckets), perk_rows, sum(v for v in buckets.values())


SCOPE_NOTE = {
    "layers": "annotated call sites (step3; SGLang = decoder layers only)",
    "forward": "whole forward incl. lm_head/sampling tail (step1)",
}


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--phase", choices=["prefill", "decode"], required=True,
                    help="labels the report only; the workbooks decide the phase")
    ap.add_argument("--src", nargs=2, action="append", metavar=("LABEL", "XLSX"),
                    required=True, help="repeatable (>=2): LABEL step1-or-step3.xlsx")
    ap.add_argument("--out", required=True, metavar="CSV")
    ap.add_argument("--top", type=int, default=6,
                    help="top-N kernels per bucket in the detail section")
    args = ap.parse_args()
    if len(args.src) < 2:
        ap.error("--src must be given at least twice")

    sides = []
    for label, path in args.src:
        scope, rows = read_workbook(path)
        buckets, perk, total = bucketize(rows)
        sides.append(dict(label=label, path=path, scope=scope, buckets=buckets,
                          perk=perk, total=total, nrows=len(rows)))
    scopes = {s["scope"] for s in sides}
    if len(scopes) > 1:
        raise SystemExit("[ERROR] mixed scopes: "
                         + ", ".join(f"{s['label']}={s['scope']}" for s in sides)
                         + " — pass the same workbook kind (step3 or step1) on every side")
    scope = scopes.pop()

    cats = sorted({c for s in sides for c in s["buckets"]},
                  key=lambda c: BUCKET_ORDER.index(c) if c in BUCKET_ORDER else 99)
    labels = [s["label"] for s in sides]

    print(f"\n=== GLM-5.2 {args.phase} buckets — scope: {SCOPE_NOTE[scope]} ===")
    for s in sides:
        print(f"  {s['label']:<12} sigma={s['total']/1000:8.2f} ms   "
              f"{s['nrows']} rows   {Path(s['path']).name}")
    w0 = 24
    print("\n  " + "Bucket".ljust(w0) + "".join(f"{l+'_ms':>12}" for l in labels)
          + (f"{'D(A-B)ms':>12}{'A/B':>8}" if len(sides) == 2 else ""))
    print("  " + "-" * (w0 + 12 * len(sides) + (20 if len(sides) == 2 else 0)))
    for c in cats:
        vals = [s["buckets"].get(c, 0.0) / 1000 for s in sides]
        line = "  " + c.ljust(w0) + "".join(f"{v:>12.2f}" for v in vals)
        if len(sides) == 2:
            a, b = vals
            line += f"{a-b:>12.2f}" + (f"{a/b:>8.2f}" if b > 1e-9 else f"{'inf':>8}")
        print(line)
    print("  " + "-" * (w0 + 12 * len(sides) + (20 if len(sides) == 2 else 0)))
    tots = [s["total"] / 1000 for s in sides]
    line = "  " + "TOTAL sigma_kernel".ljust(w0) + "".join(f"{t:>12.2f}" for t in tots)
    if len(sides) == 2:
        a, b = tots
        line += f"{a-b:>12.2f}" + (f"{a/b:>8.2f}" if b > 1e-9 else f"{'inf':>8}")
    print(line)

    # CSV layout is consumed verbatim by side_by_side.py --summary-csv.
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([f"# GLM-5.2 {args.phase} buckets"] + labels)
        w.writerow([f"# scope", SCOPE_NOTE[scope]])
        for s in sides:
            w.writerow([f"# {s['label']}", Path(s['path']).name,
                        f"sigma_ms={s['total']/1000:.2f}", f"rows={s['nrows']}"])
        w.writerow([])
        hdr = ["Bucket"] + [f"{l}_ms" for l in labels]
        if len(sides) == 2:
            hdr += [f"Delta_{labels[0]}_minus_{labels[1]}_ms",
                    f"{labels[0]}_over_{labels[1]}"]
        w.writerow(hdr)
        for c in cats:
            vals = [round(s["buckets"].get(c, 0.0) / 1000, 3) for s in sides]
            row = [c] + vals
            if len(sides) == 2:
                a, b = vals
                row += [round(a - b, 3), round(a / b, 3) if b > 1e-9 else ""]
            w.writerow(row)
        row = ["TOTAL"] + [round(t, 3) for t in tots]
        if len(sides) == 2:
            a, b = tots
            row += [round(a - b, 3), round(a / b, 3) if b > 1e-9 else ""]
        w.writerow(row)
        w.writerow([])
        w.writerow([f"# top-{args.top} kernels per bucket"])
        w.writerow(["Side", "Bucket", "Kernel", "Sum_ms", "Count"])
        for s in sides:
            by_bucket = defaultdict(list)
            for b, nm, us, c in s["perk"]:
                by_bucket[b].append((us, nm, c))
            for c in cats:
                for us, nm, cnt in sorted(by_bucket.get(c, []), reverse=True)[:args.top]:
                    w.writerow([s["label"], c, nm[:80], round(us / 1000, 3), cnt])
    print(f"\n[INFO] CSV written: {out}")


if __name__ == "__main__":
    main()
