"""Dump every kernel in a step3/step1 workbook with its call site and layer type.

Bucket rules written from kernel names alone mis-file this model: TileLang emits
`main_kernel` for whatever it compiled, so the same name can be sparse MLA on one
row and mHC on another. step3 carries LeafModule (source file + line), which
disambiguates them, and LayerType, which the analyzer already derived from the
graph-off module tree.

    python3 trace_analysis/diagnostics/dump_kernel_inventory.py <step3.xlsx> [...]
"""

import sys
from collections import defaultdict

from openpyxl import load_workbook


def dump(path):
    wb = load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}

    is_step3 = "LeafModule" in idx
    agg = defaultdict(lambda: [0.0, 0, set(), set()])
    for r in rows:
        if r is None or r[0] is None:
            continue
        try:
            if is_step3:
                name = str(r[idx["KernelName"]])
                ms = float(r[idx["SumDuration_us"]] or 0) / 1000.0
                cnt = int(float(r[idx["LaunchCount"]] or 0))
                leaf = str(r[idx["LeafModule"]])
                lt = str(r[idx["LayerType"]])
            else:
                name = str(r[idx["Name"]])
                ms = float(r[idx["SumDuration_us"]] or 0) / 1000.0
                cnt = int(float(r[idx["Count"]] or 0))
                leaf = lt = "-"
        except (KeyError, TypeError, ValueError):
            continue
        a = agg[name]
        a[0] += ms
        a[1] += cnt
        a[2].add(leaf)
        a[3].add(lt)

    print(f"===== {path.split('/')[-1]}   ({len(agg)} distinct kernels)")
    total = sum(v[0] for v in agg.values())
    for name, (ms, cnt, leaves, lts) in sorted(agg.items(), key=lambda kv: -kv[1][0]):
        pct = 100.0 * ms / total if total else 0.0
        print(f"{ms:9.3f} ms {pct:5.1f}%  n={cnt:<5d} {name[:78]}")
        if is_step3:
            print(f"{'':22}   call={','.join(sorted(leaves))[:70]}  layer={','.join(sorted(lts))[:40]}")
    print(f"{total:9.3f} ms  TOTAL\n")


if __name__ == "__main__":
    for p in sys.argv[1:]:
        dump(p)
